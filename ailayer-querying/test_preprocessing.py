"""
Unit tests for preprocessing.py's extraction/chunking logic.

Deterministic, pure-function tests — no Qdrant, no network, no real files on
disk. DOCX fixtures are built in-memory via python-docx; PDF fixtures via
PyMuPDF's own document-creation API. Covers the parts of the extraction
pipeline that don't require sentence-transformers/torch (never imported here).
"""
import io

import fitz
import pytest
from docx import Document

from preprocessing import (
    _cell,
    _is_docx,
    _serialize_table,
    chunk_text,
    clean_text,
    extract_full,
)


# ── clean_text ────────────────────────────────────────────────────────────

def test_clean_text_collapses_excess_blank_lines():
    assert clean_text("line one\n\n\n\n\nline two") == "line one\n\nline two"


def test_clean_text_collapses_repeated_spaces_and_tabs():
    assert clean_text("word1    word2\t\tword3") == "word1 word2 word3"


def test_clean_text_strips_control_characters_but_keeps_tab_and_newline():
    text = clean_text("keep\ttab\nnewline\x00\x01drop control chars")
    assert "\x00" not in text and "\x01" not in text
    assert "keep\ttab\nnewline" in text


def test_clean_text_normalizes_legal_term_case():
    result = clean_text("The PLAINTIFF and Defendant signed an AFFIDAVIT.")
    assert "plaintiff" in result
    assert "defendant" in result
    assert "affidavit" in result


def test_clean_text_trims_leading_and_trailing_whitespace():
    assert clean_text("   padded text   ") == "padded text"


# ── _cell ─────────────────────────────────────────────────────────────────

def test_cell_none_becomes_empty_string():
    assert _cell(None) == ""


def test_cell_collapses_internal_whitespace():
    assert _cell("  multi   space   value  ") == "multi space value"


def test_cell_stringifies_non_string_values():
    assert _cell(42) == "42"


# ── _serialize_table ──────────────────────────────────────────────────────

def test_serialize_table_includes_page_and_index_header():
    out = _serialize_table(["Name", "Age"], [["Alice", "30"]], page=0, t_idx=0)
    assert "[Table 1, page 1]" in out
    assert "Headers: Name | Age." in out
    assert "Row 1: Name=Alice, Age=30." in out


def test_serialize_table_skips_fully_empty_rows():
    out = _serialize_table(["A", "B"], [["", ""], ["x", "y"]], page=0, t_idx=0)
    # Row numbering reflects original position (the empty row still "counts"),
    # it's just not printed — only one "Row" line should appear, labeled 2.
    assert out.count("Row") == 1
    assert "Row 2: A=x, B=y." in out


def test_serialize_table_falls_back_to_column_numbers_when_headers_blank():
    out = _serialize_table(["", ""], [["x", "y"]], page=0, t_idx=0)
    assert "Column 1" in out and "Column 2" in out


# ── chunk_text ────────────────────────────────────────────────────────────

def test_chunk_text_empty_input_returns_no_chunks():
    assert chunk_text("") == []
    assert chunk_text("   ") == []


def test_chunk_text_short_text_returns_single_chunk():
    text = "one two three four five"
    chunks = chunk_text(text, chunk_size=200, overlap=32)
    assert chunks == [text]


def test_chunk_text_splits_long_text_into_overlapping_windows():
    words = [f"word{i}" for i in range(500)]
    text = " ".join(words)
    chunks = chunk_text(text, chunk_size=200, overlap=32)

    assert len(chunks) > 1
    # Every word appears somewhere in the chunked output (no silent data loss)
    reconstructed = " ".join(chunks).split()
    assert set(words) <= set(reconstructed)


def test_chunk_text_consecutive_chunks_actually_overlap():
    words = [f"w{i}" for i in range(250)]
    text = " ".join(words)
    chunks = chunk_text(text, chunk_size=200, overlap=32)

    assert len(chunks) >= 2
    first_words = chunks[0].split()
    second_words = chunks[1].split()
    overlap_words = set(first_words[-32:]) & set(second_words[:32])
    assert len(overlap_words) > 0


# ── _is_docx ──────────────────────────────────────────────────────────────

def test_is_docx_detects_by_filename_extension():
    assert _is_docx(b"anything", "report.docx") is True
    assert _is_docx(b"anything", "report.pdf") is False


def test_is_docx_detects_by_zip_magic_bytes_when_no_filename():
    docx_like_bytes = b"PK\x03\x04rest of zip data"
    assert _is_docx(docx_like_bytes, None) is True

    pdf_like_bytes = b"%PDF-1.4 rest of pdf data"
    assert _is_docx(pdf_like_bytes, None) is False


# ── extract_full — DOCX (built in-memory, real python-docx round-trip) ────

def _make_docx_bytes(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row_values in enumerate(table_rows):
            for c, value in enumerate(row_values):
                table.cell(r, c).text = value
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_extract_full_docx_body_text():
    docx_bytes = _make_docx_bytes(["First paragraph.", "Second paragraph."])
    result = extract_full(docx_bytes, filename="statement.docx")

    assert "First paragraph." in result.text
    assert "Second paragraph." in result.text
    assert result.tables == []
    assert result.combined == result.text


def test_extract_full_docx_with_table_produces_serialized_table_and_excludes_it_from_body():
    docx_bytes = _make_docx_bytes(
        ["Intro paragraph."],
        table_rows=[["Name", "Role"], ["Officer Rajan", "Police"]],
    )
    result = extract_full(docx_bytes, filename="report.docx")

    assert "Intro paragraph." in result.text
    assert len(result.tables) == 1
    assert result.tables[0].headers == ["Name", "Role"]
    assert result.tables[0].rows == [["Officer Rajan", "Police"]]
    assert "Officer Rajan" in result.combined
    # Table content shouldn't leak into the plain body text extraction
    assert "Officer Rajan" not in result.text


# ── extract_full — PDF (built in-memory via PyMuPDF) ───────────────────────

def _make_pdf_bytes(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def test_extract_full_pdf_body_text():
    pdf_bytes = _make_pdf_bytes("Evidence collected at the scene.")
    result = extract_full(pdf_bytes, filename="evidence.pdf")

    assert "Evidence collected at the scene." in result.text
    assert result.combined == result.text  # no tables detected in a plain-text PDF


def test_extract_full_falls_back_to_pdf_when_filename_and_magic_bytes_dont_say_docx():
    pdf_bytes = _make_pdf_bytes("Plain content.")
    result = extract_full(pdf_bytes, filename=None)
    assert "Plain content." in result.text
